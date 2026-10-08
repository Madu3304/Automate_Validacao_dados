"""Geração dos arquivos de saída (relatórios e planilha tratada)."""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .regras import ROTULOS_TIPOS, ROTULOS_TOTAIS

_COR_CABECALHO = "0B4F8C"


def _escrever(writer, aba: str, df: pd.DataFrame, estilizar: bool = True):
    aba = aba[:31]
    df.to_excel(writer, sheet_name=aba, index=False)
    ws = writer.sheets[aba]
    if estilizar:
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor=_COR_CABECALHO)
            celula.alignment = Alignment(vertical="center")
        if len(df):
            ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for i, coluna in enumerate(df.columns, start=1):
        larguras = [len(str(coluna))] + [len(str(v)) for v in df[coluna].head(500)]
        ws.column_dimensions[get_column_letter(i)].width = min(max(larguras) + 2, 70)


def _df(registros: list[dict], colunas: dict) -> pd.DataFrame:
    if not registros:
        return pd.DataFrame(columns=list(colunas.values()))
    return pd.DataFrame(registros)[list(colunas.keys())].rename(columns=colunas)


def _resumo(reg: dict, totais: dict) -> pd.DataFrame:
    linhas = [
        ("Execução", reg["id"]),
        ("Status", reg["status"]),
        ("Data", reg["criado_em"]),
        ("Usuário", reg["usuario"]),
        ("Versão das regras", reg["versao_regras"]),
        ("Explicação", reg.get("explicacao", "")),
    ]
    linhas += [(ROTULOS_TOTAIS.get(k, k), v) for k, v in totais.items()]
    return pd.DataFrame(linhas, columns=["Indicador", "Valor"])


def relatorio_conciliacao(caminho: Path, reg: dict, resultado: dict):
    divergencias = [{**d, "tipo": ROTULOS_TIPOS.get(d["tipo"], d["tipo"])} for d in resultado["divergencias"]]
    p = resultado["parametros"]
    parametros = [
        ("Coluna chave SAP", p["coluna_chave_sap"]),
        ("Coluna chave Klassmatt", p["coluna_chave_klassmatt"]),
        ("Normalização da chave", p["normalizacao_chave"]),
    ]
    parametros += [(f"Campo comparado: {c['campo']}",
                    f"{c['coluna_sap']} x {c['coluna_klassmatt']} ({c['comparacao']}, {c['criticidade']})")
                   for c in p["campos_comparados"]]
    parametros += [(f"Campo ignorado: {c['campo']}", c["motivo"]) for c in p["campos_ignorados"]]

    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        _escrever(w, "Resumo", _resumo(reg, resultado["totais"]))
        _escrever(w, "Divergencias", _df(divergencias, {
            "criticidade": "Criticidade", "tipo": "Tipo", "chave": "Código normalizado", "campo": "Campo",
            "valor_sap": "Valor SAP", "valor_klassmatt": "Valor Klassmatt", "linha_sap": "Linha SAP",
            "linha_klassmatt": "Linha Klassmatt", "mensagem": "Mensagem"}))
        _escrever(w, "Parametros", pd.DataFrame(parametros, columns=["Parâmetro", "Valor"]))


def planilha_tratada(caminho: Path, df: pd.DataFrame, aba: str):
    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        _escrever(w, aba or "Dados", df, estilizar=False)


def relatorio_verificacao(caminho: Path, reg: dict, resultado: dict):
    ocorrencias = [{**o, "tipo": ROTULOS_TIPOS.get(o["tipo"], o["tipo"])} for o in resultado["ocorrencias"]]
    alteracoes = [{**a, "critico": "Sim" if a["critico"] else "Não"} for a in resultado["alteracoes"]]
    p = resultado["parametros"]
    regras = [{"Regra": r["id"], "Descrição": r["descricao"], "Situação": "Aplicada",
               "Colunas": ", ".join(r["colunas"]), "Crítica": "Sim" if r["critica"] else "Não"}
              for r in p["regras_aplicadas"]]
    regras += [{"Regra": r["id"], "Descrição": r["descricao"], "Situação": r["motivo"], "Colunas": "", "Crítica": ""}
               for r in p["regras_ignoradas"]]

    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        _escrever(w, "Resumo", _resumo(reg, resultado["totais"]))
        _escrever(w, "Alteracoes", _df(alteracoes, {
            "linha": "Linha original", "coluna": "Coluna", "valor_original": "Valor original",
            "valor_ajustado": "Valor ajustado", "regra": "Regra", "descricao_regra": "Descrição da regra",
            "critico": "Ajuste crítico", "observacao": "Observação"}))
        _escrever(w, "Ocorrencias", _df(ocorrencias, {
            "severidade": "Severidade", "linha": "Linha original", "coluna": "Coluna", "tipo": "Tipo",
            "valor": "Valor", "mensagem": "Mensagem", "regra": "Regra"}))
        _escrever(w, "Regras", pd.DataFrame(regras, columns=["Regra", "Descrição", "Situação", "Colunas", "Crítica"]))
